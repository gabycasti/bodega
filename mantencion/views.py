from django.shortcuts import render
from .models import Mantencion
from vehiculo.models import Vehiculo
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse



# LISTADO MANTENCIÓN

def listado_mantencion(request):

    mantenciones = Mantencion.objects.select_related(
        "vehiculo"
    ).prefetch_related(
        "vehiculo__permisos_circulacion",
        "vehiculo__seguros"
    ).all()

    hoy = date.today()
    limite = hoy + timedelta(days=30)

    for p in mantenciones:

        # REVISIÓN TÉCNICA
        p.rt_proxima = False
        p.rt_vencida = False

        if p.fecha_revision_tecnica:

            if p.fecha_revision_tecnica < hoy:
                p.rt_vencida = True

            elif p.fecha_revision_tecnica <= limite:
                p.rt_proxima = True


        # GASES
        p.gases_proxima = False
        p.gases_vencida = False

        if p.fecha_gases:

            if p.fecha_gases < hoy:
                p.gases_vencida = True

            elif p.fecha_gases <= limite:
                p.gases_proxima = True


        # SEGURO
        p.seguro_proxima = False
        p.seguro_vencido = False

        seguro = p.vehiculo.seguros.first()

        if seguro and seguro.fecha_vencimiento:

            if seguro.fecha_vencimiento < hoy:
                p.seguro_vencido = True

            elif seguro.fecha_vencimiento <= limite:
                p.seguro_proxima = True


        # PERMISO DE CIRCULACIÓN
        for permiso in p.vehiculo.permisos_circulacion.all():

            permiso.pc_proximo = False
            permiso.pc_vencido = False

            if permiso.fecha_vencimiento:

                if permiso.fecha_vencimiento < hoy:
                    permiso.pc_vencido = True

                elif permiso.fecha_vencimiento <= limite:
                    permiso.pc_proximo = True


        # CAMBIO DE ACEITE
        p.aceite_1000 = False
        p.aceite_500 = False
        p.aceite_300 = False
        p.aceite_vencido = False

        if (
            p.kilometraje is not None
            and p.kilometraje_cambio_aceite is not None
        ):

            faltan = p.kilometraje_cambio_aceite - p.kilometraje

            if faltan <= 0:
                p.aceite_vencido = True

            elif faltan <= 300:
                p.aceite_300 = True

            elif faltan <= 500:
                p.aceite_500 = True

            elif faltan <= 1000:
                p.aceite_1000 = True


    return render(
        request,
        "listado_mantencion.html",
        {"mantenciones": mantenciones}
    )





# EXPORTAR MANTENCIONES A EXCEL

def exportar_mantenciones_excel(request):

    mantenciones = Mantencion.objects.select_related(
        "vehiculo"
    ).prefetch_related(
        "vehiculo__permisos_circulacion"
    ).all()

    orden_marcas = {
        "HYUNDAI": 1,
        "FUSO": 2,
        "CITROEN": 3,
        "CHEVROLET": 4,
        "JAC": 5,
        "NISSAN": 6,
    }

    orden_propietarios = {
        "TRANSPORTES CHECK SPA": 1,
        "BANCO CHILE": 2,
        "BANCO ESTADO": 3,
    }

    mantenciones = sorted(
        mantenciones,
        key=lambda p: (
            orden_marcas.get(
                (p.vehiculo.marca or "").upper().strip(),
                99
            ),
            orden_propietarios.get(
                (p.vehiculo.propietario or "").upper().strip(),
                99
            )
        )
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Mantenciones"

    # ==========================================================
    # CONFIGURACIÓN GENERAL
    # ==========================================================

    ws.sheet_view.showGridLines = False

    # ==========================================================
    # ENCABEZADOS
    # ==========================================================

    encabezados = [
        "PATENTE",
        "MARCA",
        "MODELO",
        "PROPIETARIO",
        "FECHA RT",
        "GASES",
        "CARGA",
        "PERMISO CIRCULACIÓN",
        "KM",
        "ACEITE",
        "MUNI",
        "LUGAR DE MANTENCIÓN",
    ]

    for columna, encabezado in enumerate(encabezados, start=1):
        celda = ws.cell(
            row=2,
            column=columna,
            value=encabezado
        )

        celda.fill = PatternFill(
            fill_type="solid",
            fgColor="A6A6A6"
        )

        celda.font = Font(
            color="000000",
            bold=True,
            size=8
        )

        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        celda.border = Border(
            left=Side(style="thin", color="FFFFFF"),
            right=Side(style="thin", color="FFFFFF"),
            top=Side(style="thin", color="FFFFFF"),
            bottom=Side(style="thin", color="FFFFFF")
        )

    ws.row_dimensions[2].height = 35

    # ==========================================================
    # DATOS
    # ==========================================================

    marca_anterior = None

    for numero, p in enumerate(mantenciones, start=1):

        marca_actual = (p.vehiculo.marca or "").upper().strip()

        # Insertar fila separadora si cambia la marca
        if marca_anterior is not None and marca_actual != marca_anterior:

            ws.append([""] * 12)
            fila_separador = ws.max_row

            for columna in range(1, 13):
                ws.cell(
                    row=fila_separador,
                    column=columna
                ).fill = PatternFill(
                    fill_type="solid",
                    fgColor="000000"
                )

            ws.row_dimensions[fila_separador].height = 15

        permiso = p.vehiculo.permisos_circulacion.first()

        fecha_permiso = ""
        municipalidad = ""

        if permiso:
            if permiso.fecha_vencimiento:
                fecha_permiso = permiso.fecha_vencimiento
            if permiso.municipalidad:
                municipalidad = permiso.municipalidad

        fila = [
            p.vehiculo.patente or "",
            p.vehiculo.marca or "",
            p.vehiculo.modelo or "",
            p.vehiculo.propietario or "",
            p.fecha_revision_tecnica or "",
            p.fecha_gases or "",
            f"{p.vehiculo.carga} KG" if p.vehiculo.carga else "",
            fecha_permiso,
            p.kilometraje if p.kilometraje is not None else "",
            p.kilometraje_cambio_aceite if p.kilometraje_cambio_aceite is not None else "",
            municipalidad,
            p.vehiculo.lugar_mantencion or "",
        ]

        ws.append(fila)
        marca_anterior = marca_actual

    # ==========================================================
    # COLORES Y FORMATOS DE CELDAS SEGÚN MUESTRA
    # ==========================================================

    FILL_CHECK_SPA = PatternFill(fill_type="solid", fgColor="1F497D")    # Azul Oscuro
    FILL_BANCO_CHILE = PatternFill(fill_type="solid", fgColor="00B0F0")  # Azul Claro
    FILL_BANCO_ESTADO = PatternFill(fill_type="solid", fgColor="FFC000") # Amarillo/Naranja
    FILL_VERDE = PatternFill(fill_type="solid", fgColor="92D050")        # Verde Claro
    FILL_NARANJA_MANT = PatternFill(fill_type="solid", fgColor="F8CBAD") # Naranja Mantención

    FONT_BLANCA_BOLD = Font(size=7, bold=True, color="FFFFFF")
    FONT_NEGRA_BOLD = Font(size=7, bold=True, color="000000")
    FONT_AZUL_BOLD = Font(size=7, bold=True, color="0070C0")

    borde = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000")
    )

    for fila in range(3, ws.max_row + 1):

        # Saltar la fila si es un separador negro (sin patente)
        if not ws.cell(row=fila, column=1).value:
            ws.row_dimensions[fila].height = 15
            continue

        ws.row_dimensions[fila].height = 15

        # Formato base a toda la fila de datos
        for col in range(1, 13):
            celda = ws.cell(row=fila, column=col)
            celda.border = borde
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            celda.font = FONT_NEGRA_BOLD

        # 1. PROPIETARIO (COLUMNA 4 / D)
        celda_propietario = ws.cell(row=fila, column=4)
        val_prop = str(celda_propietario.value or "").upper().strip()

        if "CHECK" in val_prop:
            celda_propietario.fill = FILL_CHECK_SPA
            celda_propietario.font = FONT_BLANCA_BOLD
        elif "BANCO CHILE" in val_prop or "BANCO DE CHILE" in val_prop:
            celda_propietario.fill = FILL_BANCO_CHILE
            celda_propietario.font = FONT_BLANCA_BOLD
        elif "BANCO ESTADO" in val_prop or "ESTADO" in val_prop:
            celda_propietario.fill = FILL_BANCO_ESTADO
            celda_propietario.font = FONT_NEGRA_BOLD

        # 2. FECHA RT (COL 5) Y GASES (COL 6)
        c_rt = ws.cell(row=fila, column=5)
        c_gases = ws.cell(row=fila, column=6)

        if c_rt.value and str(c_rt.value).strip() != "-":
            c_rt.fill = FILL_VERDE
        if c_gases.value and str(c_gases.value).strip() != "-":
            c_gases.fill = FILL_VERDE

        # 3. KM (COL 9) Y CAMBIO ACEITE (COL 10)
        celda_km = ws.cell(row=fila, column=9)
        celda_aceite = ws.cell(row=fila, column=10)

        if str(celda_km.value or "").startswith("**"):
            celda_km.fill = FILL_VERDE

        celda_aceite.font = FONT_AZUL_BOLD

        # 4. LUGAR DE MANTENCIÓN (COLUMNA 12 / L)
        celda_mant = ws.cell(row=fila, column=12)
        val_mant = str(celda_mant.value or "").upper().strip()

        if "FRENOS LENG" in val_mant:
            celda_mant.fill = FILL_VERDE
        elif any(k in val_mant for k in ["HYUNDAI", "KOVACS", "POMPEYO"]):
            celda_mant.fill = FILL_NARANJA_MANT

    # ==========================================================
    # FORMATO DE FECHAS Y NÚMEROS
    # ==========================================================

    columnas_fecha = [5, 6, 8]

    for fila in range(3, ws.max_row + 1):
        if ws.cell(row=fila, column=1).value:
            for columna in columnas_fecha:
                celda = ws.cell(row=fila, column=columna)
                if celda.value and celda.value != "-":
                    celda.number_format = "DD/MM/YYYY"

            ws.cell(row=fila, column=9).number_format = '0'
            ws.cell(row=fila, column=10).number_format = '0'

    # ==========================================================
    # ANCHO DE COLUMNAS (OPTIMIZADO)
    # ==========================================================

    anchos = {
        "A": 10,  # PATENTE
        "B": 9,   # MARCA
        "C": 9,   # MODELO
        "D": 22,  # PROPIETARIO
        "E": 11,  # FECHA RT
        "F": 11,  # GASES
        "G": 10,  # CARGA
        "H": 15,  # PERMISO CIRCULACIÓN
        "I": 9,   # KM
        "J": 9,   # ACEITE
        "K": 13,  # MUNI
        "L": 22,  # LUGAR DE MANTENCIÓN
    }

    for columna, ancho in anchos.items():
        ws.column_dimensions[columna].width = ancho

    # ==========================================================
    # CONGELAR ENCABEZADOS Y FILTROS
    # ==========================================================

    ws.freeze_panes = "A3"

    if ws.max_row >= 2:
        ws.auto_filter.ref = f"A2:L{ws.max_row}"

    # ==========================================================
    # CONFIGURACIÓN DE IMPRESIÓN
    # ==========================================================

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.50
    ws.page_margins.bottom = 0.50

    ws.print_title_rows = "1:2"

    if ws.max_row >= 2:
        ws.print_area = f"A1:L{ws.max_row}"

    # ==========================================================
    # RESPUESTA
    # ==========================================================

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="reporte_mantenciones.xlsx"'

    wb.save(response)

    return response





# AGREGAR MANTENCIÓN

def agregar_mantencion(request):

    vehiculos = Vehiculo.objects.filter(
        activo=True
    ).order_by("patente")

    if request.method == "POST":

        vehiculo_id = request.POST.get("vehiculo")

        fecha_revision_tecnica = request.POST.get(
            "fecha_revision_tecnica"
        )

        archivo_revision_tecnica = request.FILES.get(
            "archivo_revision_tecnica"
        )

        fecha_gases = request.POST.get(
            "fecha_gases"
        )

        archivo_gases = request.FILES.get(
            "archivo_gases"
        )

        kilometraje = request.POST.get(
            "kilometraje"
        )

        kilometraje_cambio_aceite = request.POST.get(
            "kilometraje_cambio_aceite"
        )

        observacion = request.POST.get(
            "observacion"
        )

        vehiculo = get_object_or_404(
            Vehiculo,
            id=vehiculo_id
        )

        Mantencion.objects.create(
            vehiculo=vehiculo,
            fecha_revision_tecnica=fecha_revision_tecnica or None,
            archivo_revision_tecnica=archivo_revision_tecnica,
            fecha_gases=fecha_gases or None,
            archivo_gases=archivo_gases,
            kilometraje=kilometraje or None,
            kilometraje_cambio_aceite=kilometraje_cambio_aceite or None,
            observacion=observacion
        )

        messages.success(
            request,
            "La mantención fue registrada correctamente."
        )

        return redirect("listado_mantencion")

    return render(
        request,
        "agregar_mantencion.html",
        {
            "vehiculos": vehiculos
        }
    )



# EDITAR MANTENCIÓN
def editar_mantencion(request, id):

    mantencion = get_object_or_404(
        Mantencion,
        id=id
    )

    vehiculos = Vehiculo.objects.filter(
        activo=True
    ).order_by("patente")

    if request.method == "POST":

        vehiculo_id = request.POST.get("vehiculo")

        fecha_revision_tecnica = request.POST.get(
            "fecha_revision_tecnica"
        )

        archivo_revision_tecnica = request.FILES.get(
            "archivo_revision_tecnica"
        )

        fecha_gases = request.POST.get(
            "fecha_gases"
        )

        archivo_gases = request.FILES.get(
            "archivo_gases"
        )

        kilometraje = request.POST.get(
            "kilometraje"
        )

        kilometraje_cambio_aceite = request.POST.get(
            "kilometraje_cambio_aceite"
        )

        observacion = request.POST.get(
            "observacion"
        )

        vehiculo = get_object_or_404(
            Vehiculo,
            id=vehiculo_id
        )

        mantencion.vehiculo = vehiculo
        mantencion.fecha_revision_tecnica = (
            fecha_revision_tecnica or None
        )
        mantencion.fecha_gases = (
            fecha_gases or None
        )
        mantencion.kilometraje = (
            kilometraje or None
        )
        mantencion.kilometraje_cambio_aceite = (
            kilometraje_cambio_aceite or None
        )
        mantencion.observacion = observacion

        # Solo reemplazar archivo si se seleccionó uno nuevo
        if archivo_revision_tecnica:
            mantencion.archivo_revision_tecnica = (
                archivo_revision_tecnica
            )

        if archivo_gases:
            mantencion.archivo_gases = (
                archivo_gases
            )

        mantencion.save()

        messages.success(
            request,
            "La mantención fue actualizada correctamente."
        )

        return redirect("listado_mantencion")

    return render(
        request,
        "editar_mantencion.html",
        {
            "mantencion": mantencion,
            "vehiculos": vehiculos
        }
    )







# ELIMINAR MANTENCIÓN

def eliminar_mantencion(request, id):

    mantencion = get_object_or_404(
        Mantencion,
        id=id
    )

    if request.method == "POST":

        mantencion.delete()

        messages.success(
            request,
            "La mantención fue eliminada correctamente."
        )

        return redirect("listado_mantencion")

    return render(
        request,
        "eliminar_mantencion.html",
        {
            "mantencion": mantencion
        }
    )